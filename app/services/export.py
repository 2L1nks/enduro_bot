"""Выгрузка данных в CSV и Excel."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.repositories import admin as admin_repo
from app.utils.logger import get_logger

log = get_logger(__name__)


async def export_csv(session: AsyncSession) -> tuple[bytes, str]:
    """CSV для Excel.

    BOM обязателен: без него русский Excel открывает файл
    в неверной кодировке и получается каша. Разделитель —
    точка с запятой по той же причине.
    """
    rows = await admin_repo.export_rows(session)
    if not rows:
        rows = [{"Записей нет": ""}]

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()), delimiter=";")
    writer.writeheader()
    writer.writerows(rows)

    data = buffer.getvalue().encode("utf-8-sig")
    name = f"enduro_{datetime.now(tz=settings.tz):%Y%m%d_%H%M}.csv"
    log.info("Экспорт CSV: %d строк", len(rows))
    return data, name


async def export_xlsx(session: AsyncSession) -> tuple[bytes, str]:
    """Excel-файл. Если openpyxl не установлен — откатываемся на CSV."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        log.warning("openpyxl не установлен, отдаю CSV")
        return await export_csv(session)

    rows = await admin_repo.export_rows(session)
    wb = Workbook()
    ws = wb.active
    ws.title = "Записи"

    if not rows:
        ws["A1"] = "Записей нет"
    else:
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row in rows:
            ws.append([row[h] for h in headers])

        # Ширина колонок по содержимому, иначе всё слипается
        for idx, header in enumerate(headers, start=1):
            longest = max(
                [len(str(header))] + [len(str(r[header])) for r in rows[:200]]
            )
            ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, 40)

        ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    name = f"enduro_{datetime.now(tz=settings.tz):%Y%m%d_%H%M}.xlsx"
    log.info("Экспорт Excel: %d строк", len(rows))
    return buffer.getvalue(), name
